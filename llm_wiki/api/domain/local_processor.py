"""Local document processor — runs extraction without S3 or Postgres.

Processes files directly from the workspace filesystem and updates SQLite.
Respects PDF_BACKEND config and optional Mistral/LibreOffice backends.
"""

import asyncio
import json
import logging
import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path

import aiosqlite

from config import settings

logger = logging.getLogger(__name__)


async def _extract_knowledge(db, doc_id: str, content: str) -> None:
    """Fire-and-forget knowledge extraction + entity graph + conflict detection. Never raises."""
    try:
        from services.extractor import is_available, extract_knowledge, store_extractions
        if not is_available():
            return
        if not content or not content.strip():
            return
        extractions = await extract_knowledge(content)
        if extractions:
            await store_extractions(db, doc_id, extractions)
            logger.info("Knowledge extraction completed for doc %s", doc_id[:8])

            # Build entity graph and detect conflicts from extraction results
            try:
                from services.graph_builder import build_entity_graph, detect_conflicts
                await build_entity_graph(db, doc_id)
                await detect_conflicts(db, doc_id)
            except Exception as e:
                logger.warning("Graph/conflict processing failed for doc %s: %s", doc_id[:8], str(e)[:200])
    except Exception as e:
        logger.warning("Knowledge extraction failed for doc %s: %s", doc_id[:8], str(e)[:200])

PDF_TYPES = frozenset({"pdf"})
SPREADSHEET_TYPES = frozenset({"xlsx", "xls"})
IMAGE_TYPES = frozenset({"png", "jpg", "jpeg", "webp", "gif"})
OFFICE_TYPES = frozenset({"pptx", "ppt", "docx", "doc"})

import re

from services.pdf_extract import extract_pdf


def _rewrite_image_paths(content: str, assets_rel_path: str) -> str:
    """Rewrite ![image](filename.png) to ![image](assets_rel_path/filename.png)."""
    def _replace(m: re.Match) -> str:
        alt = m.group(1)
        src = m.group(2)
        # Only rewrite bare filenames (not URLs or already-relative paths)
        if "/" not in src and not src.startswith("http"):
            return f"![{alt}]({assets_rel_path}/{src})"
        return m.group(0)

    return re.sub(r"!\[([^\]]*)\]\(([^)]+)\)", _replace, content)


async def process_document(db: aiosqlite.Connection, doc_id: str, workspace: Path) -> None:
    """Process a pending document: extract text, chunk, update index."""
    cursor = await db.execute(
        "SELECT id, filename, file_type, relative_path, status FROM documents WHERE id = ?",
        (doc_id,),
    )
    row = await cursor.fetchone()
    if not row:
        logger.warning("Document %s not found", doc_id[:8])
        return

    cols = [d[0] for d in cursor.description]
    doc = dict(zip(cols, row))

    if doc["status"] not in ("pending", "processing"):
        return

    file_type = doc["file_type"] or ""
    file_path = workspace / doc["relative_path"]

    if not file_path.is_file():
        await db.execute(
            "UPDATE documents SET status = 'failed', error_message = 'File not found', "
            "updated_at = datetime('now') WHERE id = ?",
            (doc_id,),
        )
        await db.commit()
        return

    await db.execute(
        "UPDATE documents SET status = 'processing', updated_at = datetime('now') WHERE id = ?",
        (doc_id,),
    )
    await db.commit()

    try:
        if file_type in PDF_TYPES:
            await _process_pdf(db, doc_id, file_path, workspace)
        elif file_type in OFFICE_TYPES:
            await _process_office(db, doc_id, file_path, workspace)
        elif file_type in SPREADSHEET_TYPES:
            await _process_spreadsheet(db, doc_id, file_path)
        elif file_type in IMAGE_TYPES:
            await _process_image(db, doc_id)
        elif file_type in ("html", "htm"):
            await _process_html(db, doc_id, file_path)
        else:
            await db.execute(
                "UPDATE documents SET status = 'ready', updated_at = datetime('now') WHERE id = ?",
                (doc_id,),
            )
            await db.commit()

            # Fire-and-forget knowledge extraction for text files
            try:
                text_content = file_path.read_text(encoding="utf-8", errors="replace")
                if text_content.strip():
                    asyncio.create_task(_extract_knowledge(db, doc_id, text_content))
            except Exception:
                pass

        logger.info("Processed %s: %s", doc["filename"], file_type)

    except Exception as e:
        error_msg = str(e)[:500]
        await db.execute(
            "UPDATE documents SET status = 'failed', error_message = ?, "
            "updated_at = datetime('now') WHERE id = ?",
            (error_msg, doc_id),
        )
        await db.commit()
        logger.error("Failed to process %s: %s", doc["filename"], e)


async def _store_chunks(db: aiosqlite.Connection, doc_id: str, chunks: list) -> None:
    """Store chunks into SQLite, replacing any existing ones.
    Optionally generates vector embeddings if embedding service is configured."""
    await db.execute("DELETE FROM document_chunks WHERE document_id = ?", (doc_id,))

    # Try to generate embeddings for all chunks
    embeddings: list[bytes | None] = [None] * len(chunks)
    try:
        from services.embedder import is_available, generate_embeddings, embedding_to_blob
        if is_available() and chunks:
            texts = [c.content for c in chunks]
            vecs = await generate_embeddings(texts)
            if vecs and len(vecs) == len(chunks):
                embeddings = [embedding_to_blob(v) for v in vecs]
                logger.info("Generated embeddings for %d chunks (doc=%s)", len(chunks), doc_id[:8])
    except Exception as e:
        logger.debug("Embedding generation skipped: %s", e)

    for i, c in enumerate(chunks):
        await db.execute(
            "INSERT INTO document_chunks (id, document_id, chunk_index, content, page, "
            "start_char, token_count, header_breadcrumb, embedding) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (str(uuid.uuid4()), doc_id, c.index, c.content, c.page,
             c.start_char, c.token_count, c.header_breadcrumb, embeddings[i]),
        )


# ── PDF extraction ────────────────────────────────────────────────────────

async def _store_page_contents(
    db: aiosqlite.Connection, doc_id: str,
    page_contents: list[tuple[int, str]], parser: str,
) -> None:
    """Store extracted pages, chunks, and update document status."""
    num_pages = len(page_contents)

    await db.execute("DELETE FROM document_pages WHERE document_id = ?", (doc_id,))
    for page_num, content in page_contents:
        await db.execute(
            "INSERT INTO document_pages (id, document_id, page, content) VALUES (?, ?, ?, ?)",
            (str(uuid.uuid4()), doc_id, page_num, content),
        )

    full_content = "\n\n---\n\n".join(md for _, md in page_contents)

    from services.chunker import chunk_pages
    chunks = chunk_pages(page_contents)
    await _store_chunks(db, doc_id, chunks)

    await db.execute(
        "UPDATE documents SET status = 'ready', content = ?, page_count = ?, "
        "parser = ?, updated_at = datetime('now') WHERE id = ?",
        (full_content, num_pages, parser, doc_id),
    )
    await db.commit()

    # Fire-and-forget knowledge extraction (don't block indexing)
    asyncio.create_task(_extract_knowledge(db, doc_id, full_content))


async def _process_pdf(db: aiosqlite.Connection, doc_id: str, file_path: Path, workspace: Path) -> None:
    """Extract PDF text and images. Uses opendataloader by default, Mistral if configured."""
    if settings.PDF_BACKEND == "mistral" and settings.MISTRAL_API_KEY:
        await _process_pdf_mistral(db, doc_id, file_path, workspace)
    else:
        # Extract images to wiki/assets/<pdf-stem>/
        pdf_stem = file_path.stem.replace(" ", "_")[:50]
        assets_dir = workspace / "wiki" / "assets" / pdf_stem
        page_contents = await asyncio.to_thread(extract_pdf, str(file_path), str(assets_dir))

        # Rewrite image references to be wiki-relative paths
        rewritten = []
        for page_num, content in page_contents:
            # Replace bare image filenames with relative path from wiki root
            content = _rewrite_image_paths(content, f"assets/{pdf_stem}")
            rewritten.append((page_num, content))

        await _store_page_contents(db, doc_id, rewritten, "opendataloader")


# ── Office processing ─────────────────────────────────────────────────────

async def _process_office(db: aiosqlite.Connection, doc_id: str, file_path: Path, workspace: Path) -> None:
    """Convert Office docs to PDF via local LibreOffice, then extract text."""
    lo = shutil.which("libreoffice") or shutil.which("soffice")
    if not lo:
        await db.execute(
            "UPDATE documents SET status = 'failed', "
            "error_message = 'LibreOffice not installed. Install it to process Office files.', "
            "updated_at = datetime('now') WHERE id = ?",
            (doc_id,),
        )
        await db.commit()
        return

    with tempfile.TemporaryDirectory() as tmpdir:
        result = await asyncio.to_thread(
            subprocess.run,
            [lo, "--headless", "--convert-to", "pdf", "--outdir", tmpdir, str(file_path)],
            capture_output=True, timeout=120,
        )
        if result.returncode != 0:
            raise RuntimeError(f"LibreOffice conversion failed: {result.stderr.decode()[:300]}")

        pdf_files = list(Path(tmpdir).glob("*.pdf"))
        if not pdf_files:
            raise RuntimeError("LibreOffice produced no PDF output")

        converted_pdf = pdf_files[0]

        # Store converted PDF in cache for the viewer
        cache_dir = workspace / ".llmwiki" / "cache" / "local" / doc_id
        cache_dir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(converted_pdf, cache_dir / "converted.pdf")

        # Extract images to wiki/assets/<file-stem>/
        file_stem = file_path.stem.replace(" ", "_")[:50]
        assets_dir = workspace / "wiki" / "assets" / file_stem
        page_contents = await asyncio.to_thread(extract_pdf, str(converted_pdf), str(assets_dir))

        rewritten = []
        for page_num, content in page_contents:
            content = _rewrite_image_paths(content, f"assets/{file_stem}")
            rewritten.append((page_num, content))

        await _store_page_contents(db, doc_id, rewritten, "libreoffice+opendataloader")


# ── Mistral OCR ───────────────────────────────────────────────────────────

async def _process_pdf_mistral(db: aiosqlite.Connection, doc_id: str, file_path: Path, workspace: Path) -> None:
    """Extract PDF via Mistral OCR API (better tables/layout, requires API key)."""
    import httpx
    import base64

    pdf_bytes = file_path.read_bytes()
    pdf_b64 = base64.b64encode(pdf_bytes).decode()

    async with httpx.AsyncClient(timeout=120) as client:
        resp = await client.post(
            "https://api.mistral.ai/v1/ocr",
            headers={"Authorization": f"Bearer {settings.MISTRAL_API_KEY}"},
            json={
                "model": "mistral-ocr-latest",
                "document": {"type": "document_url", "document_url": f"data:application/pdf;base64,{pdf_b64}"},
            },
        )
        resp.raise_for_status()
        result = resp.json()

    pages = result.get("pages", [])
    page_contents = [(i + 1, p.get("markdown", "")) for i, p in enumerate(pages)]
    await _store_page_contents(db, doc_id, page_contents, "mistral")


# ── Spreadsheet processing ────────────────────────────────────────────────

async def _process_spreadsheet(db: aiosqlite.Connection, doc_id: str, file_path: Path) -> None:
    """Extract spreadsheet data via openpyxl. Stores pages AND chunks for search."""
    from openpyxl import load_workbook

    wb = await asyncio.to_thread(load_workbook, str(file_path), read_only=True, data_only=True)

    await db.execute("DELETE FROM document_pages WHERE document_id = ?", (doc_id,))

    all_content = []
    page_contents = []
    for i, sheet_name in enumerate(wb.sheetnames, 1):
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(" | ".join(str(c) if c is not None else "" for c in row))
        content = "\n".join(rows)
        elements = json.dumps({"sheet_name": sheet_name})

        await db.execute(
            "INSERT INTO document_pages (id, document_id, page, content, elements) "
            "VALUES (?, ?, ?, ?, ?)",
            (str(uuid.uuid4()), doc_id, i, content, elements),
        )
        all_content.append(f"## {sheet_name}\n\n{content}")
        page_contents.append((i, content))

    num_sheets = len(wb.sheetnames)
    wb.close()
    full_content = "\n\n".join(all_content)

    from services.chunker import chunk_pages
    chunks = chunk_pages(page_contents)
    await _store_chunks(db, doc_id, chunks)

    await db.execute(
        "UPDATE documents SET status = 'ready', content = ?, page_count = ?, "
        "parser = 'openpyxl', updated_at = datetime('now') WHERE id = ?",
        (full_content, num_sheets, doc_id),
    )
    await db.commit()

    # Fire-and-forget knowledge extraction (don't block indexing)
    asyncio.create_task(_extract_knowledge(db, doc_id, full_content))


# ── Image / HTML processing ──────────────────────────────────────────────

async def _process_image(db: aiosqlite.Connection, doc_id: str) -> None:
    """Images are stored as-is — just mark ready."""
    await db.execute(
        "UPDATE documents SET status = 'ready', page_count = 1, "
        "parser = 'native', updated_at = datetime('now') WHERE id = ?",
        (doc_id,),
    )
    await db.commit()


async def _process_html(db: aiosqlite.Connection, doc_id: str, file_path: Path) -> None:
    """Extract HTML content via webmd parser."""
    raw_html = file_path.read_text(encoding="utf-8", errors="replace")

    try:
        from html_parser import Parser
        parser = Parser(raw_html, content_only=True)
        result = parser.parse()
        content = result.content
    except Exception:
        content = raw_html

    from services.chunker import chunk_text
    chunks = chunk_text(content)
    await _store_chunks(db, doc_id, chunks)

    await db.execute(
        "UPDATE documents SET status = 'ready', content = ?, page_count = 1, "
        "parser = 'webmd', updated_at = datetime('now') WHERE id = ?",
        (content, doc_id),
    )
    await db.commit()

    # Fire-and-forget knowledge extraction (don't block indexing)
    asyncio.create_task(_extract_knowledge(db, doc_id, content))
