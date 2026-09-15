"""PG-local document processor — pg_local 模式的文档提取处理器。

镜像 local_processor.py 的行为，但面向 PostgreSQL（jz_kb_* 表）与 asyncpg pool。
直接从 workspace 文件系统提取文本并更新索引，不依赖 S3。

来源：2026-09 修复 —— pg_local 模式下 /v1/upload 因 sqlite_db=None 必 500，
本地 OCR 代理（main OcrService 引擎2）因此整体失效。
"""

import asyncio
import json
import logging
import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path

from config import settings
from services.pdf_extract import extract_pdf
from domain.local_processor import _rewrite_image_paths

logger = logging.getLogger(__name__)

PDF_TYPES = frozenset({"pdf"})
SPREADSHEET_TYPES = frozenset({"xlsx", "xls"})
IMAGE_TYPES = frozenset({"png", "jpg", "jpeg", "webp", "gif"})
OFFICE_TYPES = frozenset({"pptx", "ppt", "docx", "doc"})


async def _store_chunks(pool, doc_id: str, chunks: list) -> None:
    """Store chunks into jz_kb_document_chunks, replacing any existing ones."""
    await pool.execute("DELETE FROM jz_kb_document_chunks WHERE document_id = $1", doc_id)
    for i, c in enumerate(chunks):
        await pool.execute(
            "INSERT INTO jz_kb_document_chunks (id, document_id, chunk_index, content, page, "
            "start_char, token_count, header_breadcrumb) VALUES ($1, $2, $3, $4, $5, $6, $7, $8)",
            str(uuid.uuid4()), doc_id, c.index, c.content, c.page,
            c.start_char, c.token_count, c.header_breadcrumb,
        )


async def _store_page_contents(pool, doc_id: str, page_contents: list[tuple[int, str]], parser: str) -> None:
    """Store extracted pages, chunks, and update document status."""
    num_pages = len(page_contents)

    await pool.execute("DELETE FROM jz_kb_document_pages WHERE document_id = $1", doc_id)
    for page_num, content in page_contents:
        await pool.execute(
            "INSERT INTO jz_kb_document_pages (id, document_id, page, content) VALUES ($1, $2, $3, $4)",
            str(uuid.uuid4()), doc_id, page_num, content,
        )

    full_content = "\n\n---\n\n".join(md for _, md in page_contents)

    from services.chunker import chunk_pages
    chunks = chunk_pages(page_contents)
    await _store_chunks(pool, doc_id, chunks)

    await pool.execute(
        "UPDATE jz_kb_documents SET status = 'ready', content = $2, page_count = $3, "
        "parser = $4, updated_at = NOW() WHERE id = $1",
        doc_id, full_content, num_pages, parser,
    )


async def _process_pdf(pool, doc_id: str, file_path: Path, workspace: Path) -> None:
    """Extract PDF text and images via opendataloader (Java) or Mistral if configured."""
    if settings.PDF_BACKEND == "mistral" and settings.MISTRAL_API_KEY:
        await _process_pdf_mistral(pool, doc_id, file_path)
        return

    pdf_stem = file_path.stem.replace(" ", "_")[:50]
    assets_dir = workspace / "wiki" / "assets" / pdf_stem
    page_contents = await asyncio.to_thread(extract_pdf, str(file_path), str(assets_dir))

    rewritten = [
        (page_num, _rewrite_image_paths(content, f"assets/{pdf_stem}"))
        for page_num, content in page_contents
    ]
    await _store_page_contents(pool, doc_id, rewritten, "opendataloader")


async def _process_pdf_mistral(pool, doc_id: str, file_path: Path) -> None:
    """Extract PDF via Mistral OCR API (requires API key)."""
    import base64
    import httpx

    pdf_b64 = base64.b64encode(file_path.read_bytes()).decode()
    async with httpx.AsyncClient(timeout=120) as client:
        resp = await client.post(
            "https://api.mistral.ai/v1/ocr",
            headers={"Authorization": f"Bearer {settings.MISTRAL_API_KEY}"},
            json={
                "model": "mistral-ocr-latest",
                "document": {"type": "document_url", "document_url": f"data:application/pdf;base64,{pdf_b64}"},
            },
        )
        resp.raise_for_status()
        result = resp.json()

    pages = result.get("pages", [])
    page_contents = [(i + 1, p.get("markdown", "")) for i, p in enumerate(pages)]
    await _store_page_contents(pool, doc_id, page_contents, "mistral")


async def _process_office(pool, doc_id: str, file_path: Path, workspace: Path) -> None:
    """Convert Office docs to PDF via local LibreOffice, then extract text."""
    lo = shutil.which("libreoffice") or shutil.which("soffice")
    if not lo:
        await pool.execute(
            "UPDATE jz_kb_documents SET status = 'failed', "
            "error_message = 'LibreOffice not installed. Install it to process Office files.', "
            "updated_at = NOW() WHERE id = $1",
            doc_id,
        )
        return

    with tempfile.TemporaryDirectory() as tmpdir:
        result = await asyncio.to_thread(
            subprocess.run,
            [lo, "--headless", "--convert-to", "pdf", "--outdir", tmpdir, str(file_path)],
            capture_output=True, timeout=120,
        )
        if result.returncode != 0:
            raise RuntimeError(f"LibreOffice conversion failed: {result.stderr.decode()[:300]}")

        pdf_files = list(Path(tmpdir).glob("*.pdf"))
        if not pdf_files:
            raise RuntimeError("LibreOffice produced no PDF output")

        converted_pdf = pdf_files[0]

        cache_dir = workspace / ".llmwiki" / "cache" / "local" / doc_id
        cache_dir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(converted_pdf, cache_dir / "converted.pdf")

        file_stem = file_path.stem.replace(" ", "_")[:50]
        assets_dir = workspace / "wiki" / "assets" / file_stem
        page_contents = await asyncio.to_thread(extract_pdf, str(converted_pdf), str(assets_dir))

        rewritten = [
            (page_num, _rewrite_image_paths(content, f"assets/{file_stem}"))
            for page_num, content in page_contents
        ]
        await _store_page_contents(pool, doc_id, rewritten, "libreoffice+opendataloader")


async def _process_spreadsheet(pool, doc_id: str, file_path: Path) -> None:
    """Extract spreadsheet data via openpyxl."""
    from openpyxl import load_workbook

    wb = await asyncio.to_thread(load_workbook, str(file_path), read_only=True, data_only=True)

    await pool.execute("DELETE FROM jz_kb_document_pages WHERE document_id = $1", doc_id)

    all_content = []
    page_contents = []
    for i, sheet_name in enumerate(wb.sheetnames, 1):
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(" | ".join(str(c) if c is not None else "" for c in row))
        content = "\n".join(rows)
        elements = json.dumps({"sheet_name": sheet_name})

        await pool.execute(
            "INSERT INTO jz_kb_document_pages (id, document_id, page, content, elements) "
            "VALUES ($1, $2, $3, $4, $5)",
            str(uuid.uuid4()), doc_id, i, content, elements,
        )
        all_content.append(f"## {sheet_name}\n\n{content}")
        page_contents.append((i, content))

    num_sheets = len(wb.sheetnames)
    wb.close()
    full_content = "\n\n".join(all_content)

    from services.chunker import chunk_pages
    chunks = chunk_pages(page_contents)
    await _store_chunks(pool, doc_id, chunks)

    await pool.execute(
        "UPDATE jz_kb_documents SET status = 'ready', content = $2, page_count = $3, "
        "parser = 'openpyxl', updated_at = NOW() WHERE id = $1",
        doc_id, full_content, num_sheets,
    )


async def _process_html(pool, doc_id: str, file_path: Path) -> None:
    """Extract HTML content via webmd parser."""
    raw_html = file_path.read_text(encoding="utf-8", errors="replace")

    try:
        from html_parser import Parser
        parser = Parser(raw_html, content_only=True)
        result = parser.parse()
        content = result.content
    except Exception:
        content = raw_html

    from services.chunker import chunk_text
    chunks = chunk_text(content)
    await _store_chunks(pool, doc_id, chunks)

    await pool.execute(
        "UPDATE jz_kb_documents SET status = 'ready', content = $2, page_count = 1, "
        "parser = 'webmd', updated_at = NOW() WHERE id = $1",
        doc_id, content,
    )


async def process_document(pool, doc_id: str, workspace: Path) -> None:
    """Process a pending document in pg_local mode: extract text, chunk, update index."""
    row = await pool.fetchrow(
        "SELECT id, filename, file_type, relative_path, status FROM jz_kb_documents WHERE id = $1",
        doc_id,
    )
    if not row:
        logger.warning("Document %s not found", doc_id[:8])
        return

    doc = dict(row)
    if doc["status"] not in ("pending", "processing"):
        return

    file_type = doc["file_type"] or ""
    file_path = workspace / doc["relative_path"]

    if not file_path.is_file():
        await pool.execute(
            "UPDATE jz_kb_documents SET status = 'failed', error_message = 'File not found', "
            "updated_at = NOW() WHERE id = $1",
            doc_id,
        )
        return

    await pool.execute(
        "UPDATE jz_kb_documents SET status = 'processing', updated_at = NOW() WHERE id = $1",
        doc_id,
    )

    try:
        if file_type in PDF_TYPES:
            await _process_pdf(pool, doc_id, file_path, workspace)
        elif file_type in OFFICE_TYPES:
            await _process_office(pool, doc_id, file_path, workspace)
        elif file_type in SPREADSHEET_TYPES:
            await _process_spreadsheet(pool, doc_id, file_path)
        elif file_type in IMAGE_TYPES:
            await pool.execute(
                "UPDATE jz_kb_documents SET status = 'ready', page_count = 1, "
                "parser = 'native', updated_at = NOW() WHERE id = $1",
                doc_id,
            )
        elif file_type in ("html", "htm"):
            await _process_html(pool, doc_id, file_path)
        else:
            await pool.execute(
                "UPDATE jz_kb_documents SET status = 'ready', updated_at = NOW() WHERE id = $1",
                doc_id,
            )

        logger.info("Processed (pg_local) %s: %s", doc["filename"], file_type)

    except Exception as e:
        error_msg = str(e)[:500]
        await pool.execute(
            "UPDATE jz_kb_documents SET status = 'failed', error_message = $2, "
            "updated_at = NOW() WHERE id = $1",
            doc_id, error_msg,
        )
        logger.error("Failed to process (pg_local) %s: %s", doc["filename"], e)
